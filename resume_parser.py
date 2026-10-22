import os
import json
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
from pypdf import PdfReader
from docx import Document

load_dotenv()
my_api_key=os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API key kaha hai bhai")

client=Groq(api_key=my_api_key)
model = os.getenv("GROQ_MODEL", "openai/gpt-oss-120b")


job_description="""
Description
Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.
Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.
Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older
Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills
"""


class JobD(BaseModel):
    role: str | None = None
    required_skills: list[str] = []
    preferred_skills: list[str] = []
    minimum_experience: float | None = None
    education_requirements: list[str] = []
    responsibilities: list[str] = []


class MatchResult(BaseModel):
    candidate_name: str | None = None
    score: float = 0
    matching_skills: list[str] = []
    missing_skills: list[str] = []
    experience_met: bool | None = None
    verdict: str = ""


class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []


class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


jobd_schema = JobD.model_json_schema()
resume_schema = Resume.model_json_schema()
match_schema = MatchResult.model_json_schema()


def ask_llm_json(messages):
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    return json.loads(response.choices[0].message.content)


def parse_job(job_text):
    system_prompt = f"""
    You are an expert HR assistant.

    Your job is to analyze job descriptions and extract
    structured information from them.

    Return ONLY valid JSON matching this schema:

    {jobd_schema}
    IMPORTANT:
    Do NOT return the schema itself.
    Do NOT return fields like "properties", "title" or "type".
    Fill the schema with actual information extracted from the job description.

    If minimum experience is not mentioned, return null.
    If information for a list is missing, return an empty list.
    Do not invent information.
    """
    user_prompt = f"""
    Analyze the following job description:

    {job_text}
    """
    messages=[
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]
    return JobD(**ask_llm_json(messages))


def final_score(job, resume):
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}

    Return ONLY valid JSON matching this schema
    (fill it with values, do NOT return the schema itself):

    {match_schema}

    Field meanings:
    - candidate_name: the candidate's name
    - score: overall match percentage from 0 to 100
    - matching_skills: job skills the candidate has
    - missing_skills: important job skills the candidate is missing
    - experience_met: whether the experience requirement is met
    - verdict: a short final verdict (1-2 sentences)
    """
    messages=[{"role": "user", "content": prompt}]
    result = MatchResult(**ask_llm_json(messages))
    # some models answer 0-1 instead of 0-100
    if 0 < result.score <= 1:
        result.score *= 100
    result.score = max(0, min(100, round(result.score, 1)))
    return result


def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    6. projects, education and certifications must be lists of plain strings.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    messages=[
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]
    return Resume(**ask_llm_json(messages))


# file can be a path or an open file object (e.g. an upload)
def read_pdf(file):
    reader = PdfReader(file)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text


def read_docx(file):
    document = Document(file)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file, filename):
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = read_pdf(file)
    elif suffix == ".docx":
        text = read_docx(file)
    else:
        raise ValueError(f"Unsupported file type: {suffix} (use .pdf or .docx)")
    if not text.strip():
        raise ValueError("No text found in file (is it a scanned image?)")
    return text


def print_candidate(candidate):
    print(candidate["name"], "-", candidate["score"], "%")
    print("  Matching:", ", ".join(candidate["matching_skills"]))
    print("  Missing:", ", ".join(candidate["missing_skills"]))
    print("  Verdict:", candidate["verdict"])


def main():
    job = parse_job(job_description)
    print("Role:", job.role)

    resume_folder = Path("resumes")
    all_results=[]
    for file_path in resume_folder.iterdir():
        if file_path.suffix.lower() not in [".pdf", ".docx"]:
            continue
        print("\nProcessing:", file_path.name)
        try:
            resume_text = read_resume(file_path, file_path.name)
            parsed_resume=parse_resume(resume_text) # llm call1
            time.sleep(5)  # stay under the groq rate limit
            result = final_score(job, parsed_resume) #llm call2
            time.sleep(5)
        except Exception as error:
            print("  Skipped:", error)
            continue
        print("Score:", result.score)
        all_results.append({
            "name": parsed_resume.name or result.candidate_name or file_path.stem,
            **result.model_dump(exclude={"candidate_name"}),
        })

    all_results.sort(key=lambda candidate: candidate["score"], reverse=True)

    # with fewer than 4 resumes, top and bottom would overlap
    half = len(all_results) // 2
    top_2 = all_results[:min(2, len(all_results) - half)]
    worst_2 = all_results[-min(2, half):] if half else []

    print("\nTOP CANDIDATES")
    for candidate in top_2:
        print_candidate(candidate)

    print("\nLOWEST CANDIDATES")
    for candidate in worst_2:
        print_candidate(candidate)


if __name__ == "__main__":
    main()
